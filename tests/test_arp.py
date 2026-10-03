import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import Workbook

from pint_live.arp import ArpEntry, ArpLoadError, ArpTable, load_arp_xlsx, load_arp_xlsx_many


class ArpMergeTests(unittest.TestCase):
    def test_duplicate_macs_keep_latest_ip_and_last_nonblank_hostname(self):
        table = ArpTable(entries=[
            ArpEntry("192.0.2.1", "AA:BB:CC:DD:EE:FF", "printer"),
            ArpEntry("192.0.2.2", "aabb.ccdd.eeff"),
        ])
        self.assertEqual(len(table), 1)
        self.assertEqual(table.lookup("aa-bb-cc-dd-ee-ff").ip, "192.0.2.2")
        self.assertEqual(table.lookup("aabbccddeeff").hostname, "printer")

    def test_repeated_merges_deduplicate_entries_and_sources_without_mutation(self):
        source = ArpTable(
            source_paths=[Path("arp.xlsx")],
            entries=[ArpEntry("192.0.2.1", "aabbccddeeff", "printer")],
        )
        merged = ArpTable()
        merged.extend(source)
        merged.extend(source)
        merged.extend(ArpTable(entries=[ArpEntry("192.0.2.2", "aabbccddeeff")]))
        self.assertEqual(len(merged), 1)
        self.assertEqual(merged.file_count, 1)
        self.assertEqual(merged.entries[0].ip, "192.0.2.2")
        self.assertEqual(merged.entries[0].hostname, "printer")
        self.assertEqual(source.entries[0].ip, "192.0.2.1")
        merged.extend(ArpTable(entries=[ArpEntry("192.0.2.3", "aabbccddeeff", "new")]))
        self.assertEqual(len(merged), 1)
        self.assertEqual(merged.entries[0].hostname, "new")

    def test_resolution_preserves_unmatched_mac_positions(self):
        table = ArpTable(entries=[ArpEntry("192.0.2.1", "aabbccddeeff", "printer")])
        macs = ["001122334455", "aabb.ccdd.eeff"]
        self.assertEqual(table.resolve_ips(macs), ["", "192.0.2.1"])
        self.assertEqual(table.resolve_hostnames(macs), ["", "printer"])


class ArpImportTests(unittest.TestCase):
    def test_real_workbooks_merge_while_invalid_file_is_reported(self):
        with tempfile.TemporaryDirectory() as tmp:
            paths = [Path(tmp) / name for name in ("first.xlsx", "bad.xlsx", "last.xlsx")]
            rows = [
                [["IP Address", "MAC Address", "Hostname"], ["192.0.2.1", "AA:BB:CC:DD:EE:FF", "printer"]],
                [["Wrong header"], ["bad"]],
                [["IP", "MAC"], ["192.0.2.2", "aabb.ccdd.eeff"]],
            ]
            for path, workbook_rows in zip(paths, rows):
                wb = Workbook()
                for row in workbook_rows:
                    wb.active.append(row)
                wb.save(path)
                wb.close()
            table, errors = load_arp_xlsx_many(paths)
        self.assertEqual(len(errors), 1)
        self.assertEqual(errors[0][0], paths[1])
        self.assertEqual(len(table), 1)
        self.assertEqual(table.file_count, 2)
        self.assertEqual(table.entries[0].ip, "192.0.2.2")
        self.assertEqual(table.entries[0].hostname, "printer")

    def test_validation_failures_always_close_workbook(self):
        for rows in ([], [("Wrong header",)]):
            with self.subTest(rows=rows):
                wb = Mock()
                wb.active.iter_rows.return_value = iter(rows)
                with patch("pint_live.arp.load_workbook", return_value=wb):
                    with self.assertRaises(ArpLoadError):
                        load_arp_xlsx(Path("invalid.xlsx"))
                wb.close.assert_called_once()

    def test_lazy_read_failure_is_reported_and_next_file_still_loads(self):
        def broken_rows():
            yield ("IP", "MAC")
            raise ValueError("Malformed worksheet row")

        bad = Mock()
        bad.active.iter_rows.return_value = broken_rows()
        good = Mock()
        good.active.iter_rows.return_value = iter([
            ("IP", "MAC"), ("192.0.2.1", "aabbccddeeff"),
        ])
        with patch("pint_live.arp.load_workbook", side_effect=[bad, good]):
            table, errors = load_arp_xlsx_many(["bad.xlsx", "good.xlsx"])
        self.assertEqual(len(table), 1)
        self.assertEqual(len(errors), 1)
        self.assertIn("Malformed worksheet row", errors[0][1])
        bad.close.assert_called_once()
        good.close.assert_called_once()


if __name__ == "__main__":
    unittest.main()
