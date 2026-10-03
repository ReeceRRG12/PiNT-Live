"""Workbook round-trip checks for literal data and worksheet navigation."""

from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from pint_live.arp import ArpEntry, ArpTable
from pint_live.exporters.excel import export
from pint_live.models import InterfaceEntry, LagEntry, MacEntry, NeighborEntry, ParsedSwitchData


def _interface(**overrides) -> InterfaceEntry:
    values = dict(
        port="1/1/1", link="Up", state="Forward", duplex="Full",
        speed="1G", tag="No", pvid="1",
    )
    values.update(overrides)
    return InterfaceEntry(**values)


def _linked_sheet(workbook, cell):
    target = cell.hyperlink.target
    name = target.removeprefix("#'").removesuffix("'!A1").replace("''", "'")
    return workbook[name]


class ExcelExportTests(unittest.TestCase):
    def test_sheet_names_are_valid_unique_and_navigation_resolves(self):
        names = [
            "Summary", "summary", "Site:West/[]?*\\", "'quoted'", "O'Brien",
            "x" * 40, "x" * 39 + "y", "Core", "Core Raw",
        ]
        devices = [
            ParsedSwitchData(
                host=f"10.0.0.{index}", hostname=name,
                interfaces=[_interface()], lags=[LagEntry(lag_id="1")],
                raw_outputs={"show version": "version"},
            )
            for index, name in enumerate(names, start=1)
        ]
        with tempfile.TemporaryDirectory() as directory:
            output = export(devices, Path(directory) / "navigation.xlsx", include_raw_outputs=True)
            workbook = load_workbook(output)
            try:
                self.assertEqual(workbook.sheetnames[0], "Summary")
                self.assertEqual(len(workbook.sheetnames), 1 + 3 * len(devices))
                self.assertEqual(len({name.casefold() for name in workbook.sheetnames}), len(workbook.sheetnames))
                for name in workbook.sheetnames:
                    self.assertLessEqual(len(name), 31)
                    self.assertNotRegex(name, r"[\\/*?:\[\]]")
                    self.assertFalse(name.startswith("'") or name.endswith("'"))
                summary = workbook["Summary"]
                for row, device in enumerate(devices, start=2):
                    self.assertEqual(summary.cell(row, 1).value, device.hostname)
                    for column, suffix in [(9, ""), (10, " Raw"), (11, " LAGs")]:
                        sheet = _linked_sheet(workbook, summary.cell(row, column))
                        self.assertNotEqual(sheet.title, "Summary")
                        self.assertTrue(sheet.title.endswith(suffix))
                        self.assertEqual(sheet["A1"].hyperlink.target, "#'Summary'!A1")
                    main = _linked_sheet(workbook, summary.cell(row, 9))
                    self.assertEqual(main["A5"].value, "1/1/1")
                    self.assertEqual(main.auto_filter.ref, "A4:N5")
            finally:
                workbook.close()

    def test_device_and_arp_strings_remain_literal_and_counts_remain_numeric(self):
        device = ParsedSwitchData(
            host="10.0.0.1", hostname="=switch", model="=model", firmware="=firmware",
            interfaces=[_interface(description="=1+1", state="#N/A")],
            mac_table=[MacEntry(mac="0011.2233.4455", port="1/1/1", vlan="1", entry_type="Dynamic")],
            neighbors=[NeighborEntry(
                protocol="=protocol", local_port="1/1/1", device_id="=device",
                management_ip="=address", remote_port="=port", platform="=platform",
            )],
            lags=[LagEntry(lag_id="1", name="=lag", mode="=mode", interface="=interface")],
            raw_outputs={"=command": "=raw-output\n#N/A"},
        )
        arp = ArpTable(entries=[ArpEntry(ip="10.0.0.5", mac="001122334455", hostname="=arp-host")])
        with tempfile.TemporaryDirectory() as directory:
            output = export([device], Path(directory) / "literal.xlsx", arp_table=arp, include_raw_outputs=True)
            workbook = load_workbook(output, data_only=False)
            try:
                summary = workbook["Summary"]
                main = _linked_sheet(workbook, summary["I2"])
                lag = _linked_sheet(workbook, summary["K2"])
                raw = _linked_sheet(workbook, summary["J2"])
                for sheet, address, expected in [
                    (summary, "A2", "=switch"), (summary, "C2", "=model"),
                    (summary, "D2", "=firmware"), (main, "C5", "#N/A"),
                    (main, "J5", "=arp-host"), (main, "K5", "=protocol"),
                    (main, "L5", "=device"), (main, "M5", "=address"),
                    (main, "N5", "=port"), (main, "O5", "=platform"),
                    (main, "P5", "=1+1"), (lag, "B3", "=lag"),
                    (lag, "C3", "=mode"), (lag, "E3", "=interface"),
                    (raw, "A5", "=command"), (raw, "C5", "=raw-output"),
                    (raw, "C6", "#N/A"),
                ]:
                    with self.subTest(sheet=sheet.title, address=address):
                        self.assertEqual(sheet[address].value, expected)
                        self.assertEqual(sheet[address].data_type, "s")
                for address, expected in [("E2", 1), ("F2", 1), ("G2", 0), ("H2", 0)]:
                    self.assertEqual(summary[address].value, expected)
                    self.assertEqual(summary[address].data_type, "n")
                self.assertEqual(raw["B5"].data_type, "n")
                self.assertEqual(main.auto_filter.ref, "A4:P5")
                self.assertEqual(main["A2"].data_type, "s")
            finally:
                workbook.close()

    def test_illegal_controls_are_removed_without_losing_tabs_and_newlines(self):
        device = ParsedSwitchData(
            host="10.0.0.1", hostname="SW\x08-01", model="model\x00name",
            interfaces=[_interface(description="rack\x08name\tline\nnext")],
            raw_outputs={"show\x00 version": "version\x08text"},
        )
        with tempfile.TemporaryDirectory() as directory:
            output = export([device], Path(directory) / "controls.xlsx", include_raw_outputs=True)
            workbook = load_workbook(output)
            try:
                self.assertEqual(workbook["Summary"]["A2"].value, "SW-01")
                self.assertEqual(workbook["SW-01"]["N5"].value, "rackname\tline\nnext")
                self.assertIn("modelname", workbook["SW-01"]["A2"].value)
                self.assertEqual(workbook["SW-01 Raw"]["A5"].value, "show version")
                self.assertEqual(workbook["SW-01 Raw"]["C5"].value, "versiontext")
            finally:
                workbook.close()

    def test_empty_export_still_has_a_summary(self):
        with tempfile.TemporaryDirectory() as directory:
            output = export([], Path(directory) / "empty.xlsx")
            workbook = load_workbook(output)
            try:
                self.assertEqual(workbook.sheetnames, ["Summary"])
                self.assertEqual(workbook["Summary"]["A1"].value, "Switch")
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
